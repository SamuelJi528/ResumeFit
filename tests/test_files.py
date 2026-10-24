from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from file_utils import extract_text, MAX_FILE_BYTES


def test_utf8_resume_text_keeps_names():
    assert extract_text("resume.txt", "\ufeff名字\nReact developer".encode("utf-8")) == "名字\nReact developer"


def test_docx_import_includes_table_content():
    document = Document()
    document.add_paragraph("Built a React interface.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Python inventory project"
    output = BytesIO()
    document.save(output)
    text = extract_text("resume.docx", output.getvalue())
    assert "React" in text and "Python" in text


def test_pdf_import_reads_selectable_text():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Built a React interface.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    assert "Built a React interface." in extract_text("resume.pdf", output.getvalue())


def test_blank_pdf_reports_missing_text():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="No selectable text"):
        extract_text("resume.pdf", output.getvalue())


@pytest.mark.parametrize("name,data", [("resume.exe", b"text"), ("resume.docx", b"bad file"),
                                       ("resume.txt", b"x" * (MAX_FILE_BYTES + 1)), ("resume.txt", b"")],
                         ids=["unsupported-extension", "corrupt-docx", "oversized-file", "empty-file"])
def test_invalid_files_fail_cleanly(name, data):
    with pytest.raises(ValueError):
        extract_text(name, data)
