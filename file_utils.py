from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


MAX_FILE_BYTES = 3 * 1024 * 1024


def extract_text(name: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Use a file smaller than 3 MB.")
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".txt":
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ValueError("TXT files must use UTF-8. Paste the text instead.") from None
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Use an unlocked PDF, or paste the resume text.")
            if len(reader.pages) > 12:
                raise ValueError("Upload a resume with 12 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("This DOCX expands to too much data. Paste the resume text instead.")
            document = Document(BytesIO(data))
            lines = [p.text for p in document.paragraphs]
            lines.extend(" | ".join(cell.text for cell in row.cells)
                         for table in document.tables for row in table.rows)
            text = "\n".join(lines)
        else:
            raise ValueError("Upload a PDF, DOCX, or UTF-8 TXT file.")
    except ValueError:
        raise
    except Exception:
        raise ValueError("This file could not be read. Try another file or paste its text.") from None
    if not text.strip():
        raise ValueError("No selectable text was found. Scanned PDFs need OCR; paste the text instead.")
    return text.strip()
